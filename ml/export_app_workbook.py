#!/usr/bin/env python3
"""Build the supervisor-facing Excel workbook for deployed app classifiers."""

from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

NAVY = "1F4E79"
BLUE = "2E75B6"
LIGHT = "D6EAF8"
PALE = "F2F2F2"
AMBER = "FFF2CC"
WHITE = "FFFFFF"
GREEN = "E2EFDA"
THIN = Border(
    left=Side(style="thin", color="BFBFBF"),
    right=Side(style="thin", color="BFBFBF"),
    top=Side(style="thin", color="BFBFBF"),
    bottom=Side(style="thin", color="BFBFBF"),
)
TITLE_FONT = Font(name="Calibri", size=16, bold=True, color=WHITE)
HEADER_FONT = Font(name="Calibri", size=11, bold=True, color=WHITE)
SECTION_FONT = Font(name="Calibri", size=12, bold=True, color=NAVY)
BODY_FONT = Font(name="Calibri", size=11)
NOTE_FONT = Font(name="Calibri", size=10, italic=True, color="595959")
NA_FONT = Font(name="Calibri", size=11, italic=True, color="808080")


def _fill(color: str) -> PatternFill:
    return PatternFill("solid", fgColor=color)


def _style_header_row(ws: Worksheet, row: int, start: int, end: int) -> None:
    for col in range(start, end + 1):
        cell = ws.cell(row, col)
        cell.fill = _fill(NAVY)
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = THIN


def _style_range(ws: Worksheet, min_row: int, max_row: int, min_col: int, max_col: int, number_format: str | None = None) -> None:
    for row in range(min_row, max_row + 1):
        for col in range(min_col, max_col + 1):
            cell = ws.cell(row, col)
            cell.font = BODY_FONT
            cell.border = THIN
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            if row % 2 == 0:
                cell.fill = _fill(PALE)
            if number_format and isinstance(cell.value, (int, float)):
                cell.number_format = number_format


def _set_widths(ws: Worksheet, widths: dict[int, float]) -> None:
    for col, width in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = width


def _write_title(ws: Worksheet, title: str, cols: int) -> None:
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=cols)
    cell = ws.cell(1, 1, title)
    cell.fill = _fill(NAVY)
    cell.font = TITLE_FONT
    cell.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 24


def _put(ws: Worksheet, row: int, col: int, value, *, bold=False, fill=None, fmt=None, font=None):
    cell = ws.cell(row, col, value)
    cell.font = font or (Font(name="Calibri", size=11, bold=bold) if bold else BODY_FONT)
    cell.border = THIN
    cell.alignment = Alignment(vertical="center", wrap_text=True)
    if fill:
        cell.fill = _fill(fill)
    if fmt and isinstance(value, (int, float)):
        cell.number_format = fmt
    return cell


def _history_stats(history: pd.DataFrame) -> dict:
    df = history.copy()
    df["epoch"] = range(1, len(df) + 1)
    last = df.iloc[-1]
    best_val_acc_i = int(df["val_accuracy"].idxmax())
    best_train_acc_i = int(df["accuracy"].idxmax())
    best_val_loss_i = int(df["val_loss"].idxmin())
    return {
        "n_epochs": int(len(df)),
        "final_epoch": int(last["epoch"]),
        "final_train_acc": float(last["accuracy"]),
        "final_val_acc": float(last["val_accuracy"]),
        "final_train_loss": float(last["loss"]),
        "final_val_loss": float(last["val_loss"]),
        "best_train_acc": float(df.loc[best_train_acc_i, "accuracy"]),
        "best_train_acc_epoch": int(df.loc[best_train_acc_i, "epoch"]),
        "best_val_acc": float(df.loc[best_val_acc_i, "val_accuracy"]),
        "best_val_acc_epoch": int(df.loc[best_val_acc_i, "epoch"]),
        "best_val_loss": float(df.loc[best_val_loss_i, "val_loss"]),
        "best_val_loss_epoch": int(df.loc[best_val_loss_i, "epoch"]),
    }


def _metric_block(metrics: dict) -> list[tuple[str, object]]:
    return [
        ("n", metrics["n"]),
        ("Accuracy", metrics["accuracy"]),
        ("Balanced accuracy", metrics["balanced_accuracy"]),
        ("Precision (macro)", metrics["precision_macro"]),
        ("Recall (macro)", metrics["recall_macro"]),
        ("F1 (macro)", metrics["f1_macro"]),
        ("Precision (weighted)", metrics["precision_weighted"]),
        ("Recall (weighted)", metrics["recall_weighted"]),
        ("F1 (weighted)", metrics["f1_weighted"]),
        ("ROC AUC (Effective)", metrics["roc_auc"]),
        ("Average precision (Effective)", metrics["average_precision"]),
        ("Support Ineffective", metrics["support"]["Ineffective"]),
        ("Support Effective", metrics["support"]["Effective"]),
    ]


def write_model_summary(ws: Worksheet, binary: dict, rgb3: dict, history_stats: dict | None) -> None:
    _write_title(ws, "PhageScan — deployed app classifiers", 8)
    ws.merge_cells("A2:H2")
    ws["A2"] = (
        "Only the two methods used by the current mobile app. "
        "Binary CNN metrics are RGB-label agreement, not independently verified biological accuracy. "
        "3-class RGB accuracy metrics are N/A because no independent 3-class labels exist."
    )
    ws["A2"].font = NOTE_FONT
    ws["A2"].alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[2].height = 42

    headers = [
        "Method",
        "Classification type",
        "Classes",
        "Sample counts",
        "Accuracy",
        "Precision / Recall / F1 (macro)",
        "ROC AUC",
        "Evaluation limitation",
    ]
    for i, header in enumerate(headers, 1):
        ws.cell(4, i, header)
    _style_header_row(ws, 4, 1, 8)

    demo = binary["metrics_full_demo_vs_csv_labels"]
    test = binary.get("original_split_metrics", {}).get("test")
    sample_counts = "Demo 94 wells (47 Effective / 47 Ineffective)"
    if test:
        sample_counts += f"; original split train/val/test = 56/19/19 (held-out test 12 Ineffective / 7 Effective)"
    hist = ""
    if history_stats:
        hist = (
            f" CNN_2 log: final train/val acc {history_stats['final_train_acc']:.1%}/"
            f"{history_stats['final_val_acc']:.1%}; first perfect val acc at epoch "
            f"{history_stats['best_val_acc_epoch']}."
        )

    cnn_row = [
        "Binary CNN (well_classifier_cnn.h5)",
        "Supervised CNN (machine learning)",
        "Ineffective, Effective",
        sample_counts,
        demo["accuracy"],
        f"{demo['precision_macro']:.4f} / {demo['recall_macro']:.4f} / {demo['f1_macro']:.4f}",
        demo["roc_auc"],
        "100% demo and 100% original held-out test agree with RGB-generated binary labels. "
        "Not wet-lab verified." + hist,
    ]
    rgb_counts = rgb3["prediction_counts_demo"]
    n = rgb3["n_demo"]
    rgb_row = [
        "Deterministic 3-class RGB rule",
        "Rule-based (not trained)",
        "Effective, Somewhat Effective, Ineffective",
        f"Demo {n} wells — predicted {rgb_counts.get('Effective', 0)} Effective / "
        f"{rgb_counts.get('Somewhat Effective', 0)} Somewhat Effective / "
        f"{rgb_counts.get('Ineffective', 0)} Ineffective",
        "N/A",
        "N/A",
        "N/A",
        rgb3["reason_metrics_omitted"],
    ]
    for r_i, values in enumerate((cnn_row, rgb_row), 5):
        fill = LIGHT if r_i == 5 else GREEN
        for c_i, value in enumerate(values, 1):
            cell = _put(ws, r_i, c_i, value, fill=fill)
            if c_i in (5, 7) and isinstance(value, float):
                cell.number_format = "0.00%"
            if value == "N/A":
                cell.font = NA_FONT
                cell.fill = _fill(AMBER)
        ws.row_dimensions[r_i].height = 72

    ws.merge_cells("A8:H10")
    ws["A8"] = (
        "Do not compare Binary CNN accuracy with 3-class RGB output counts. They are different tasks. "
        "The app Binary button uses the CNN. The app Multi-Class button uses this RGB rule (`model_type=rule`). "
        "The API 4-class CNN and historical Random Forest / 3-class CNN are excluded from this workbook."
    )
    ws["A8"].font = NOTE_FONT
    ws["A8"].alignment = Alignment(wrap_text=True, vertical="top")
    ws["A8"].fill = _fill(AMBER)
    _set_widths(ws, {1: 36, 2: 28, 3: 36, 4: 42, 5: 14, 6: 32, 7: 12, 8: 55})
    ws.freeze_panes = "A5"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def write_history(ws: Worksheet, history: pd.DataFrame, stats: dict) -> None:
    _write_title(ws, "Binary CNN training history (CNN_2.ipynb log — not a retrain)", 6)
    ws.merge_cells("A2:F2")
    ws["A2"] = (
        "30 epochs from CNN_2.csv. Architecture matches the deployed softmax binary CNN. "
        "Distinguish final-epoch scores (epoch 30) from the first epoch that reached the best validation accuracy."
    )
    ws["A2"].font = NOTE_FONT
    ws["A2"].alignment = Alignment(wrap_text=True)
    ws.row_dimensions[2].height = 32

    headers = ["Epoch", "Training accuracy", "Validation accuracy", "Training loss", "Validation loss"]
    for i, header in enumerate(headers, 1):
        ws.cell(4, i, header)
    _style_header_row(ws, 4, 1, 5)

    df = history.copy()
    df.insert(0, "epoch", range(1, len(df) + 1))
    for offset, row in enumerate(df.itertuples(index=False), 5):
        ws.cell(offset, 1, int(row.epoch))
        ws.cell(offset, 2, float(row.accuracy))
        ws.cell(offset, 3, float(row.val_accuracy))
        ws.cell(offset, 4, float(row.loss))
        ws.cell(offset, 5, float(row.val_loss))
    last_data = 4 + len(df)
    _style_range(ws, 5, last_data, 1, 5)
    for row in range(5, last_data + 1):
        ws.cell(row, 2).number_format = "0.00%"
        ws.cell(row, 3).number_format = "0.00%"
        ws.cell(row, 4).number_format = "0.0000"
        ws.cell(row, 5).number_format = "0.0000"

    ws.cell(4, 7, "Summary statistic")
    ws.cell(4, 8, "Value")
    _style_header_row(ws, 4, 7, 8)
    summary_rows = [
        ("Epochs", stats["n_epochs"], "0"),
        ("Final epoch", stats["final_epoch"], "0"),
        ("Final training accuracy", stats["final_train_acc"], "0.00%"),
        ("Final validation accuracy", stats["final_val_acc"], "0.00%"),
        ("Final training loss", stats["final_train_loss"], "0.0000"),
        ("Final validation loss", stats["final_val_loss"], "0.0000"),
        ("Best training accuracy", stats["best_train_acc"], "0.00%"),
        ("Epoch of best training accuracy", stats["best_train_acc_epoch"], "0"),
        ("Best validation accuracy", stats["best_val_acc"], "0.00%"),
        ("Epoch of best validation accuracy (first)", stats["best_val_acc_epoch"], "0"),
        ("Best (lowest) validation loss", stats["best_val_loss"], "0.000000"),
        ("Epoch of best validation loss", stats["best_val_loss_epoch"], "0"),
    ]
    for i, (name, value, fmt) in enumerate(summary_rows, 5):
        _put(ws, i, 7, name, fill=LIGHT)
        cell = _put(ws, i, 8, value)
        cell.number_format = fmt

    acc = LineChart()
    acc.title = "Training and validation accuracy"
    acc.style = 10
    acc.y_axis.title = "Accuracy"
    acc.x_axis.title = "Epoch"
    acc.y_axis.scaling.min = 0
    acc.y_axis.scaling.max = 1.05
    acc.height = 8
    acc.width = 15
    data = Reference(ws, min_col=2, min_row=4, max_col=3, max_row=last_data)
    cats = Reference(ws, min_col=1, min_row=5, max_row=last_data)
    acc.add_data(data, titles_from_data=True)
    acc.set_categories(cats)
    acc.shape = 4
    ws.add_chart(acc, "A37")

    loss = LineChart()
    loss.title = "Training and validation loss"
    loss.style = 12
    loss.y_axis.title = "Loss"
    loss.x_axis.title = "Epoch"
    loss.height = 8
    loss.width = 15
    data = Reference(ws, min_col=4, min_row=4, max_col=5, max_row=last_data)
    loss.add_data(data, titles_from_data=True)
    loss.set_categories(cats)
    ws.add_chart(loss, "E37")

    _set_widths(ws, {1: 12, 2: 22, 3: 22, 4: 18, 5: 18, 6: 4, 7: 42, 8: 16})
    ws.freeze_panes = "A5"


def write_metrics(ws: Worksheet, binary: dict) -> None:
    _write_title(ws, "Binary CNN evaluation metrics", 8)
    ws.merge_cells("A2:H2")
    ws["A2"] = (
        "Held-out test (19 original CNN_2 wells) is separate from the current 94-well demo plate. "
        "Both scores measure agreement with RGB-generated binary labels (distance < 20 → Effective), "
        "not independently verified biological effectiveness. API prediction mapping was verified."
    )
    ws["A2"].font = NOTE_FONT
    ws["A2"].alignment = Alignment(wrap_text=True)
    ws.row_dimensions[2].height = 36

    sections = [
        (
            4,
            "Original held-out test (19 images) vs serving-aligned RGB rule",
            binary.get("original_split_metrics", {}).get("test"),
            "These 19 wells were held out of the CNN_2 60/20/20 split (random_state=42). Support: 12 Ineffective, 7 Effective.",
        ),
        (
            22,
            "Full current demo (94 images) vs data/ground_truth_labels.csv",
            binary["metrics_full_demo_vs_csv_labels"],
            "Current demo crops are not the original training files (no shared MD5 hashes). Labels match the binary RGB rule on this plate.",
        ),
    ]
    if binary.get("original_split_metrics", {}).get("train"):
        sections.append(
            (
                40,
                "Original training set (56 images) vs RGB rule — not a test score",
                binary["original_split_metrics"]["train"],
                "Training-set agreement with the RGB rule. Do not report this as held-out performance.",
            )
        )
        sections.append(
            (
                58,
                "Original validation set (19 images) vs RGB rule — not the held-out test",
                binary["original_split_metrics"]["val"],
                "Validation-set agreement with the RGB rule during the original split. Separate from the 19-image test set.",
            )
        )

    for start, title, metrics, note in sections:
        if metrics is None:
            ws.cell(start, 1, title)
            ws.cell(start, 1).font = SECTION_FONT
            ws.cell(start + 1, 1, "Not recovered in this run.")
            continue
        ws.merge_cells(start_row=start, start_column=1, end_row=start, end_column=6)
        ws.cell(start, 1, title).font = SECTION_FONT
        ws.merge_cells(start_row=start + 1, start_column=1, end_row=start + 1, end_column=6)
        ws.cell(start + 1, 1, note).font = NOTE_FONT

        ws.cell(start + 3, 1, "Metric")
        ws.cell(start + 3, 2, "Value")
        _style_header_row(ws, start + 3, 1, 2)
        for i, (name, value) in enumerate(_metric_block(metrics), start + 4):
            _put(ws, i, 1, name, fill=LIGHT)
            cell = _put(ws, i, 2, value)
            if isinstance(value, float) and name not in {"n"}:
                cell.number_format = "0.00%" if value <= 1.0000001 else "0.00"
            if name.startswith("Support") or name == "n":
                cell.number_format = "0"

        ws.cell(start + 3, 4, "Class")
        ws.cell(start + 3, 5, "Support")
        ws.cell(start + 3, 6, "Precision")
        ws.cell(start + 3, 7, "Recall")
        ws.cell(start + 3, 8, "F1")
        _style_header_row(ws, start + 3, 4, 8)
        for i, cls in enumerate(["Ineffective", "Effective"]):
            r = start + 4 + i
            _put(ws, r, 4, cls, fill=LIGHT)
            _put(ws, r, 5, metrics["support"][cls], fmt="0")
            _put(ws, r, 6, metrics["precision"][cls], fmt="0.00%")
            _put(ws, r, 7, metrics["recall"][cls], fmt="0.00%")
            _put(ws, r, 8, metrics["f1"][cls], fmt="0.00%")
        _put(ws, start + 6, 4, "Macro", fill=LIGHT, bold=True)
        _put(ws, start + 6, 5, metrics["n"], fmt="0")
        _put(ws, start + 6, 6, metrics["precision_macro"], fmt="0.00%")
        _put(ws, start + 6, 7, metrics["recall_macro"], fmt="0.00%")
        _put(ws, start + 6, 8, metrics["f1_macro"], fmt="0.00%")
        _put(ws, start + 7, 4, "Weighted", fill=LIGHT, bold=True)
        _put(ws, start + 7, 5, metrics["n"], fmt="0")
        _put(ws, start + 7, 6, metrics["precision_weighted"], fmt="0.00%")
        _put(ws, start + 7, 7, metrics["recall_weighted"], fmt="0.00%")
        _put(ws, start + 7, 8, metrics["f1_weighted"], fmt="0.00%")

    _set_widths(ws, {1: 38, 2: 16, 3: 4, 4: 16, 5: 12, 6: 14, 7: 12, 8: 12})
    ws.freeze_panes = "A4"


def write_confusion(ws: Worksheet, binary: dict) -> None:
    _write_title(ws, "Binary CNN confusion matrices", 8)
    ws.merge_cells("A2:H2")
    ws["A2"] = "Rows = true label. Columns = predicted label. Normalized matrices are row-normalized (recall view)."
    ws["A2"].font = NOTE_FONT

    blocks = [
        (4, "Original held-out test (19) — raw counts", binary.get("original_split_metrics", {}).get("test"), False),
        (12, "Original held-out test (19) — row-normalized", binary.get("original_split_metrics", {}).get("test"), True),
        (20, "Full demo (94) — raw counts", binary["metrics_full_demo_vs_csv_labels"], False),
        (28, "Full demo (94) — row-normalized", binary["metrics_full_demo_vs_csv_labels"], True),
    ]
    for start, title, metrics, normalized in blocks:
        ws.merge_cells(start_row=start, start_column=1, end_row=start, end_column=4)
        ws.cell(start, 1, title).font = SECTION_FONT
        if metrics is None:
            ws.cell(start + 1, 1, "Not recovered.")
            continue
        labels = metrics["confusion_matrix"]["labels"]
        matrix = metrics["confusion_matrix"]["normalized_true"] if normalized else metrics["confusion_matrix"]["counts"]
        ws.cell(start + 1, 1, "true \\ pred")
        for j, name in enumerate(labels):
            _put(ws, start + 1, j + 2, f"Predicted {name}", bold=True, fill=NAVY, font=HEADER_FONT)
        _put(ws, start + 1, 1, "true \\ pred", bold=True, fill=NAVY, font=HEADER_FONT)
        for i, name in enumerate(labels):
            _put(ws, start + 2 + i, 1, f"True {name}", bold=True, fill=LIGHT)
            for j, value in enumerate(matrix[i]):
                fmt = "0.00%" if normalized else "0"
                _put(ws, start + 2 + i, j + 2, float(value) if normalized else int(value), fmt=fmt)

    _set_widths(ws, {1: 42, 2: 22, 3: 22, 4: 18})


def write_rgb(ws: Worksheet, predictions: pd.DataFrame, rgb3: dict) -> None:
    _write_title(ws, "Deterministic 3-class RGB classifier — current demo plate", 6)
    ws.merge_cells("A2:F2")
    ws["A2"] = (
        "Production rule from api/classification_utils.py: distance < 10 → Effective; "
        "≥10 and <25 → Somewhat Effective; ≥25 → Ineffective. "
        "This method has no training epochs, loss, accuracy, precision, recall, F1, or ROC AUC."
    )
    ws["A2"].font = NOTE_FONT
    ws["A2"].alignment = Alignment(wrap_text=True)
    ws.row_dimensions[2].height = 36

    ws.cell(4, 1, "Methodology")
    ws.cell(4, 1).font = SECTION_FONT
    method_rows = [
        ("Type", "Deterministic rule-based classifier (not a trained model)"),
        ("Feature", "RGB-peak Euclidean distance after dropping histogram bin 0"),
        ("Effective", "distance < 10"),
        ("Somewhat Effective", "10 ≤ distance < 25"),
        ("Ineffective", "distance ≥ 25"),
        ("Accuracy / precision / recall / F1 / ROC AUC", "N/A — no independent 3-class reference labels"),
    ]
    ws.cell(5, 1, "Item")
    ws.cell(5, 2, "Value")
    _style_header_row(ws, 5, 1, 2)
    for i, (name, value) in enumerate(method_rows, 6):
        _put(ws, i, 1, name, fill=LIGHT)
        cell = _put(ws, i, 2, value)
        if str(value).startswith("N/A"):
            cell.font = NA_FONT
            cell.fill = _fill(AMBER)

    counts = rgb3["prediction_counts_demo"]
    n = int(rgb3["n_demo"])
    ws.cell(13, 1, "Predicted class totals (94 demo wells)")
    ws.cell(13, 1).font = SECTION_FONT
    for i, header in enumerate(["Class", "Count", "Percent of demo wells"], 1):
        ws.cell(14, i, header)
    _style_header_row(ws, 14, 1, 3)
    for i, cls in enumerate(["Effective", "Somewhat Effective", "Ineffective"]):
        _put(ws, 15 + i, 1, cls, fill=LIGHT)
        _put(ws, 15 + i, 2, counts.get(cls, 0), fmt="0")
        _put(ws, 15 + i, 3, counts.get(cls, 0) / n if n else 0, fmt="0.00%")
    _put(ws, 18, 1, "Total", bold=True, fill=LIGHT)
    _put(ws, 18, 2, n, fmt="0", bold=True)
    _put(ws, 18, 3, 1, fmt="0.00%", bold=True)

    ws.cell(20, 1, "RGB distance summary by predicted class")
    ws.cell(20, 1).font = SECTION_FONT
    for i, header in enumerate(["Predicted class", "n", "Mean", "Std. dev.", "Minimum", "Maximum"], 1):
        ws.cell(21, i, header)
    _style_header_row(ws, 21, 1, 6)
    grouped = predictions.groupby("rgb_3class_predicted")["rgb_distance"]
    for i, cls in enumerate(["Effective", "Somewhat Effective", "Ineffective"]):
        series = grouped.get_group(cls) if cls in grouped.groups else pd.Series(dtype=float)
        _put(ws, 22 + i, 1, cls, fill=LIGHT)
        _put(ws, 22 + i, 2, int(series.size), fmt="0")
        _put(ws, 22 + i, 3, float(series.mean()) if len(series) else "N/A", fmt="0.00")
        _put(ws, 22 + i, 4, float(series.std(ddof=1)) if len(series) > 1 else "N/A", fmt="0.00")
        _put(ws, 22 + i, 5, float(series.min()) if len(series) else "N/A", fmt="0.00")
        _put(ws, 22 + i, 6, float(series.max()) if len(series) else "N/A", fmt="0.00")

    ws.cell(26, 1, "Per-image predictions")
    ws.cell(26, 1).font = SECTION_FONT
    headers = ["Filename", "RGB distance", "Predicted class", "Effective if distance <", "Somewhat Effective if distance <"]
    for i, header in enumerate(headers, 1):
        ws.cell(27, i, header)
    _style_header_row(ws, 27, 1, 5)
    pred = predictions.sort_values(["rgb_3class_predicted", "Filename"]).reset_index(drop=True)
    for i, row in pred.iterrows():
        r = 28 + i
        _put(ws, r, 1, row["Filename"])
        _put(ws, r, 2, float(row["rgb_distance"]), fmt="0.00")
        _put(ws, r, 3, row["rgb_3class_predicted"])
        _put(ws, r, 4, int(row.get("effective_if_distance_lt", 10)))
        _put(ws, r, 5, int(row.get("somewhat_effective_if_distance_lt", 25)))
        if r % 2 == 0:
            for col in range(1, 6):
                if ws.cell(r, col).fill.fgColor is None or ws.cell(r, col).fill.fgColor.rgb == "00000000":
                    ws.cell(r, col).fill = _fill(PALE)

    _set_widths(ws, {1: 28, 2: 16, 3: 22, 4: 26, 5: 34, 6: 14})
    ws.freeze_panes = "A28"
    ws.auto_filter.ref = f"A27:E{27 + len(pred)}"


def write_methods(ws: Worksheet, binary: dict, rgb3: dict, history_stats: dict | None) -> None:
    _write_title(ws, "Methods, data, and limitations", 2)
    ws.merge_cells("A2:B2")
    ws["A2"] = "Use this sheet when drafting manuscript methods. Nothing here is independently verified wet-lab accuracy."
    ws["A2"].font = NOTE_FONT

    paragraphs = [
        (
            "Deployed methods",
            "1) Binary CNN: api/models/well_classifier_cnn.h5, app button Analyze Image (Binary), API model_type=binary. "
            "Input 128×128, /255.0, OpenCV BGR via api.classification_utils.preprocess_image. "
            "Output softmax over BINARY_LABELS = [Ineffective, Effective] (argmax 0/1). "
            "Architecture: Conv2D(16)–MaxPool–Conv2D(32)–MaxPool–Flatten–Dense(64)–Dropout(0.3)–Dense(2, softmax). "
            "2) Deterministic 3-class RGB rule: api.classification_utils._classify_well_rgb, app button Analyze Image (Multi-Class) sends model_type=rule. "
            "Effective if RGB-peak distance < 10; Somewhat Effective if < 25; otherwise Ineffective. No learned weights.",
        ),
        (
            "What was excluded",
            "4-class CNN (API-only, not used by the app), historical 3-class CNN, Random Forest, and the binary RGB threshold-20 rule as a standalone app method. "
            "Those artifacts were not written into this workbook.",
        ),
        (
            "Datasets and sample sizes",
            "Current demo: 94 PNGs in ml/demo/segmented_wells/ (96-well plate minus circle_1_11 and circle_1_12) with ml/data/ground_truth_labels.csv "
            "(47 Effective / 47 Ineffective). Original CNN_2 training plate: 94 SegmentedWells crops used by CNN_2.ipynb. "
            "Those original files are a different image set from the current demo (zero shared MD5 hashes).",
        ),
        (
            "Train / validation / test split",
            "CNN_2.ipynb used sklearn train_test_split twice: test_size=0.4 then 0.5 on the remainder, random_state=42, no stratification. "
            "That yields 56 train / 19 validation / 19 held-out test. The 19-image test support (12 Ineffective / 7 Effective) matches the original notebook classification report. "
            "This split cannot be applied to the current demo filenames.",
        ),
        (
            "Label provenance",
            "CNN_2 training strings used an inverted RGB rule, then LabelEncoder + API mapping (argmax 1 → Effective) so serving polarity matches "
            "distance < 20 → Effective. Current ground_truth_labels.csv matches that serving-aligned binary RGB rule on the demo plate. "
            "No wet-lab plaque-assay labels are recorded in this repository. "
            "No independent 3-class labels exist; historical 3-class Actual values were generated by the same RGB rule.",
        ),
        (
            "What the numbers mean",
            "Binary CNN 100% on the 19 held-out original wells = agreement with the RGB rule on wells not used in that training split. "
            "Binary CNN 100% on the 94 demo wells = agreement with ground_truth_labels.csv (which themselves match the RGB rule). "
            "These are not independently verified biological classification accuracies. "
            "3-class RGB accuracy/precision/recall/F1/AUC are N/A because there is no independent 3-class reference.",
        ),
        (
            "API verification",
            f"Batched binary CNN predictions were checked against classification_utils._predict_label for sample wells: "
            f"{'PASS' if binary.get('api_verification', {}).get('all_match') else 'FAIL'}. "
            f"Checked files: {', '.join(binary.get('api_verification', {}).get('checked', []))}.",
        ),
        (
            "Training history correspondence",
            "CNN_2.csv is the 30-epoch Keras history saved by CNN_2.ipynb, which also saved well_classifier_cnn.h5 "
            "with this same Sequential softmax-2 architecture. The deployed file is api/models/well_classifier_cnn.h5. "
            "Final-epoch accuracy is the last row of that log. Best validation accuracy is the first epoch that reached the maximum val_accuracy. "
            "This workbook does not retrain the model.",
        ),
        (
            "Missing evidence",
            "Independently verified experimental labels; 3-class reference labels; a training notebook for any 4-class model. "
            "Do not invent those labels or fill N/A metrics with placeholder accuracies.",
        ),
    ]
    if history_stats:
        paragraphs.append(
            (
                "History snapshot for the manuscript",
                f"30 epochs. Final train/val accuracy {history_stats['final_train_acc']:.4f} / {history_stats['final_val_acc']:.4f}. "
                f"Best validation accuracy {history_stats['best_val_acc']:.4f} first reached at epoch {history_stats['best_val_acc_epoch']}. "
                f"Best validation loss {history_stats['best_val_loss']:.6g} at epoch {history_stats['best_val_loss_epoch']}.",
            )
        )

    row = 4
    for title, text in paragraphs:
        _put(ws, row, 1, title, bold=True, fill=LIGHT)
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=2)
        cell = _put(ws, row, 2, text)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[row].height = 78
        row += 1
    _set_widths(ws, {1: 32, 2: 110})


def export_workbook(path: Path, binary: dict, rgb3: dict, history: pd.DataFrame, rgb_predictions: pd.DataFrame) -> Path:
    stats = _history_stats(history) if len(history) else None
    wb = Workbook()

    ws = wb.active
    ws.title = "Model Summary"
    write_model_summary(ws, binary, rgb3, stats)

    ws = wb.create_sheet("Binary CNN History")
    write_history(ws, history, stats or {})

    ws = wb.create_sheet("Binary CNN Metrics")
    write_metrics(ws, binary)

    ws = wb.create_sheet("Binary CNN Confusion")
    write_confusion(ws, binary)

    ws = wb.create_sheet("3-Class RGB Results")
    write_rgb(ws, rgb_predictions, rgb3)

    ws = wb.create_sheet("Methods and Limitations")
    write_methods(ws, binary, rgb3, stats)

    for sheet in wb.worksheets:
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.paperSize = sheet.PAPERSIZE_TABLOID
        sheet.oddHeader.left.text = "PhageScan model evaluation — deployed app methods only"
        sheet.oddFooter.left.text = "RGB-label agreement ≠ biological ground truth"
        sheet.oddFooter.right.text = "Page &P of &N"

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path
